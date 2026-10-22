# 本项目维护的文档写入 MCP 服务；把 create_docx 注册为可发现工具。
# 它会创建目录并写入指定 DOCX，和 File 子 Agent 使用的只读文档解析工具不是同一个服务。
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from mcp.server.mcpserver import MCPServer


server = MCPServer(
    "office-writer",
    instructions=(
        "Create valid Microsoft Word .docx files. Paths must be absolute and "
        "must use the .docx extension."
    ),
)
# MCPServer 保存工具注册信息，并通过 stdio 与 MCP Client 交换请求和响应。


# 接收模型传入的输出路径，检查绝对路径及 .docx 扩展名，并创建父目录。
# 注意 mkdir 是实际副作用；此函数不是单纯的字符串检查，也没有执行文件覆盖确认。
def _validated_output_path(output_path: str) -> Path:
    """校验 DOCX 输出路径，并确保目标父目录存在。"""
    path = Path(output_path).expanduser()
    if not path.is_absolute():
        raise ValueError("output_path must be an absolute path")
    if path.suffix.lower() != ".docx":
        raise ValueError("output_path must end with .docx")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


# 输入已保存的 Path；检查 ZIP 完整性和必要成员，再用 python-docx 重读。
# 失败时抛异常，不会把“保存成功”当作文档格式已验证。
def _verify_docx(path: Path) -> None:
    """同时从 ZIP 包结构和 python-docx 解析两个层面验证生成文件。"""
    try:
        # DOCX 本质是包含多个 XML 文件的 Office Open XML ZIP 包。
        with ZipFile(path, "r") as package:
            corrupt_member = package.testzip()
            if corrupt_member:
                raise ValueError(f"corrupt OOXML member: {corrupt_member}")
            members = set(package.namelist())
    except BadZipFile as exc:
        raise ValueError("generated file is not a valid OOXML ZIP package") from exc

    # 这两个成员是一个可识别 Word 文档最基本的包文件。
    required_members = {"[Content_Types].xml", "word/document.xml"}
    missing = required_members - members
    if missing:
        raise ValueError(f"generated DOCX is missing: {sorted(missing)}")

    # Parsing it again catches malformed document XML or package relationships.
    Document(path)


# output_path/title 必填，其余列表可省略；逐项建立文档，保存并验证后返回路径说明。
# 该工具会真实写文件，已有同名文件可能被覆盖，因此没有列入 File 子 Agent 白名单。
@server.tool(
    description=(
        "Create a real Microsoft Word DOCX document with a title, paragraphs, "
        "numbered list, and bullet list. Returns the verified absolute path."
    )
)
def create_docx(
    output_path: str,
    title: str,
    paragraphs: list[str] | None = None,
    numbered_list: list[str] | None = None,
    bullet_list: list[str] | None = None,
) -> str:
    """根据标题、段落和列表创建 DOCX，验证成功后返回绝对路径。"""
    # @server.tool 装饰器把普通 Python 函数注册为可被 MCP Client 发现的工具。
    path = _validated_output_path(output_path)
    document = Document()
    document.add_heading(title, level=0)

    # 可选列表为 None 时按空列表处理；不同循环分别使用正文、有序和无序样式，而不是拼成一段纯文本。
    for paragraph in paragraphs or []:
        document.add_paragraph(paragraph)
    for item in numbered_list or []:
        document.add_paragraph(item, style="List Number")
    for item in bullet_list or []:
        document.add_paragraph(item, style="List Bullet")

    # python-docx 会生成真正的 OOXML 文件，而不是把纯文本伪装成 .docx。
    document.save(path)
    _verify_docx(path)
    return f"DOCX created and verified: {path}"


if __name__ == "__main__":
    # 仅直接运行本文件时启动 stdio Server；被导入时只完成工具定义。
    server.run(transport="stdio")
